import json
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Faculty lists from reference course_faculty.json
FACS = {
    'CSE299': ['afe', 'itn', 'mle', 'msrb', 'muo', 'nlh', 'oisd', 'smsl', 'sfr1', 'tns1'],
    'CSE311': ['fth', 'nlh', 'rih', 'tns1']
}

# Load existing detailed reviews to get exact ratings assigned previously where applicable
def load_detailed_ratings(json_path, course):
    detailed_map = {}
    try:
        with open(json_path, encoding='utf-8') as f:
            d = json.load(f)
            course_data = d.get(course.lower(), {})
            for fac, rev_list in course_data.items():
                for r in rev_list:
                    txt = r['review'].strip()
                    detailed_map[(fac.lower(), txt)] = r['rating']
    except Exception as e:
        print(f"Note: Could not load detailed ratings from {json_path}: {e}")
    return detailed_map

detailed_299 = load_detailed_ratings('cse299_detailed_reviews.json', 'CSE299')
detailed_311 = load_detailed_ratings('cse311_detailed_reviews.json', 'CSE311')

def judge_comment_rating(text, fac, detailed_map):
    norm_txt = text.strip()
    
    # 1. Check exact match in detailed reviews
    if (fac.lower(), norm_txt) in detailed_map:
        return detailed_map[(fac.lower(), norm_txt)]
    
    # Check partial match
    for (f, t), r in detailed_map.items():
        if f == fac.lower() and (t in norm_txt or norm_txt in t):
            return r

    t_lower = norm_txt.lower()

    # Drop criteria (tagging only, bump, exchange, scheduling, pure noise)
    # Check if comment is just a name tag or single word/symbol like '.', 'up', 'bump', 'cfbr', 'F', '@...'
    cleaned = re.sub(r'@[A-Za-z0-9._\s]+', '', norm_txt).strip()
    words = [w for w in re.split(r'\s+', t_lower) if w]
    
    if len(words) == 0 or norm_txt in ['.', 'F', 'f', 'up', 'bump', 'cfbr', 'thanks', 'thank you']:
        return "drop (not a review)"
    
    # Check if it's purely a tag without opinion
    if all(w.startswith('@') or w in [fac.lower(), 'sir', 'maam', 'mam', 'vaia', 'bhaiya', 'bhai'] for w in words):
        return "drop (not a review)"

    # Check logistics / exchange / section match
    if any(k in t_lower for k in ["exchange", "same section", "which section", "routine", "lab conflict", "lab section"]):
        return "drop (not a review)"

    # Heuristic ratings if it's a review
    if any(k in t_lower for k in ["goat", "11/10", "10/10", "best best", "a+ e dibe", "highly recommended", "must to do", "100 100", "top tier"]):
        return "Outstanding"
    elif any(k in t_lower for k in ["avoid", "drop", "worst", "bekar"]):
        return "avoid"
    elif any(k in t_lower for k in ["strict", "unpredictable", "tough", "insult", "mood bujha", "hard", "pera", "pressure"]):
        return "harsh"
    elif any(k in t_lower for k in ["normal", "average", "regular update", "deadline", "fair", "ok", "okay", "moderate"]):
        return "normal"
    elif any(k in t_lower for k in ["best", "valo", "bhalo", "good", "easy going", "friendly", "helpful", "curved", "go for", "recommended"]):
        return "great"

    # Default rating fallback for non-noise
    return "normal"


def extract_course_records(course, raw_file, fac_list, detailed_map):
    with open(raw_file, encoding='utf-8') as f:
        posts = json.load(f)

    records = []
    seen = set()

    for p in posts:
        ptext = p.get('text', '').strip()
        # Find which reference faculties are mentioned in post
        post_facs = [f for f in fac_list if re.search(r'\b' + re.escape(f) + r'\b', ptext, re.I)]

        # 1. Check if post text itself is a detailed review
        if post_facs and ("#facultyreview" in ptext.lower() or "review" in ptext.lower() or len(ptext) > 100):
            for fac in post_facs:
                key = (course, fac, ptext)
                if key not in seen:
                    seen.add(key)
                    rating = judge_comment_rating(ptext, fac, detailed_map)
                    records.append({
                        'course': course,
                        'faculty': fac.lower(),
                        'comment': ptext,
                        'rating': rating,
                        'source': 'Post Text'
                    })

        # 2. Check comments
        for c in p.get('comments', []):
            ctext = c.get('text', '').strip()
            if not ctext:
                continue

            comm_facs = [f for f in fac_list if re.search(r'\b' + re.escape(f) + r'\b', ctext, re.I)]
            target_facs = list(set(comm_facs if comm_facs else post_facs))

            if not target_facs:
                continue

            for fac in target_facs:
                key = (course, fac, ctext)
                if key not in seen:
                    seen.add(key)
                    rating = judge_comment_rating(ctext, fac, detailed_map)
                    records.append({
                        'course': course,
                        'faculty': fac.lower(),
                        'comment': ctext,
                        'rating': rating,
                        'source': 'Comment'
                    })

    return records

records_299 = extract_course_records('CSE299', 'raw/q_cse299.json', FACS['CSE299'], detailed_299)
records_311 = extract_course_records('CSE311', 'raw/q_cse311.json', FACS['CSE311'], detailed_311)

all_records = records_299 + records_311
print(f"Total extracted records for evaluation: {len(all_records)} (CSE299: {len(records_299)}, CSE311: {len(records_311)})")

# Build Excel File using openpyxl
wb = openpyxl.Workbook()
# Default sheet
ws_all = wb.active
ws_all.title = "All Reviews (CSE299 & 311)"

ws_299 = wb.create_sheet(title="CSE299")
ws_311 = wb.create_sheet(title="CSE311")

headers = [
    "Course",
    "Faculty",
    "review/comment",
    "ratings(given by you)",
    "User Judgment (Right/Wrong)",
    "User Remarks / Feedback"
]

# Formatting styles
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def format_sheet(ws, records):
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    for r_idx, rec in enumerate(records, start=2):
        row_data = [
            rec['course'],
            rec['faculty'].upper(),
            rec['comment'],
            rec['rating'],
            "", # User Judgment
            ""  # User Remarks
        ]
        ws.append(row_data)
        
        ws.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r_idx, column=2).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r_idx, column=3).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.cell(row=r_idx, column=4).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r_idx, column=5).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r_idx, column=6).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        for c_idx in range(1, len(headers) + 1):
            ws.cell(row=r_idx, column=c_idx).border = border_thin

    # Set specific column widths
    ws.column_dimensions['A'].width = 12 # Course
    ws.column_dimensions['B'].width = 12 # Faculty
    ws.column_dimensions['C'].width = 75 # review/comment
    ws.column_dimensions['D'].width = 22 # ratings(given by you)
    ws.column_dimensions['E'].width = 25 # User Judgment (Right/Wrong)
    ws.column_dimensions['F'].width = 40 # User Remarks / Feedback

format_sheet(ws_all, all_records)
format_sheet(ws_299, records_299)
format_sheet(ws_311, records_311)

output_excel_path = "CSE299_CSE311_Review_Evaluations.xlsx"
wb.save(output_excel_path)
print(f"Saved Excel file to {output_excel_path}")
